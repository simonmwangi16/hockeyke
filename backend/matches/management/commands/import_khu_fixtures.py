import openpyxl
from datetime import datetime, time

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Q

from competitions.models import Season, League, LeagueSeason
from teams.models import Team
from matches.models import Match


class Command(BaseCommand):
    help = "Import KHU fixtures from Excel file"

    def add_arguments(self, parser):
        parser.add_argument("file_path", type=str)
        parser.add_argument("--season", type=str, required=True)

    def clean_value(self, value):
        if value is None:
            return ""
        return str(value).strip()

    def get_team(self, value, row_number, side, league):
        team_value = self.clean_value(value)

        teams = Team.objects.filter(
            Q(name__iexact=team_value) | Q(short_name__iexact=team_value)
        )

        if league.gender:
            teams = teams.filter(gender=league.gender)

        if teams.count() == 1:
            return teams.first()

        if teams.count() == 0:
            all_matches = Team.objects.filter(
                Q(name__iexact=team_value) | Q(short_name__iexact=team_value)
            )

            debug_matches = ", ".join(
                [
                    f"{team.name} ({team.short_name}, {team.gender})"
                    for team in all_matches
                ]
            ) or "No partial DB matches"

            raise CommandError(
                f"Row {row_number}: {side} team not found: '{team_value}'. "
                f"League={league.short_name}, League gender={league.gender}. "
                f"DB matches before gender filter: {debug_matches}"
            )

        raise CommandError(
            f"Row {row_number}: Multiple {side} teams found for '{team_value}' "
            f"in gender {league.gender}: "
            + ", ".join([f"{team.name} ({team.short_name})" for team in teams])
        )

    def get_league(self, league_code, row_number):
        league_code = self.clean_value(league_code)

        try:
            return League.objects.get(short_name__iexact=league_code)
        except League.DoesNotExist:
            raise CommandError(
                f"Row {row_number}: League not found for short_name '{league_code}'"
            )
        except League.MultipleObjectsReturned:
            raise CommandError(
                f"Row {row_number}: Multiple leagues found for short_name '{league_code}'"
            )

    def normalise_date_and_status(self, value, row_number):
        if value is None or str(value).strip() == "":
            return None, "SCHEDULED"

        value_text = str(value).strip()

        if value_text.lower() == "postponed - new date tba":
            return None, "POSTPONED"

        if isinstance(value, datetime):
            return value.date(), "SCHEDULED"

        if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
            return value, "SCHEDULED"

        raise CommandError(f"Row {row_number}: Invalid match date: {value}")

    def normalise_time(self, value):
        if value in [None, ""]:
            return None

        if isinstance(value, datetime):
            return value.time()

        if isinstance(value, time):
            return value

        return value

    def handle(self, *args, **options):
        file_path = options["file_path"]
        season_name = options["season"]

        try:
            season = Season.objects.get(name=season_name)
        except Season.DoesNotExist:
            raise CommandError(f"Season not found: {season_name}")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        headers = [self.clean_value(cell.value) for cell in ws[1]]

        required_headers = [
            "DATE",
            "TIME",
            "LEAGUE",
            "MATCH NO.",
            "HOME",
            "AWAY",
            "VENUE",
        ]

        for header in required_headers:
            if header not in headers:
                raise CommandError(f"Missing required column: {header}")

        col = {header: headers.index(header) + 1 for header in required_headers}

        created_count = 0
        updated_count = 0
        skipped_count = 0

        with transaction.atomic():
            for row_number in range(2, ws.max_row + 1):
                match_date_raw = ws.cell(row_number, col["DATE"]).value
                match_time_raw = ws.cell(row_number, col["TIME"]).value
                league_code = self.clean_value(ws.cell(row_number, col["LEAGUE"]).value)
                match_number = self.clean_value(ws.cell(row_number, col["MATCH NO."]).value)
                home_name = self.clean_value(ws.cell(row_number, col["HOME"]).value)
                away_name = self.clean_value(ws.cell(row_number, col["AWAY"]).value)
                venue = self.clean_value(ws.cell(row_number, col["VENUE"]).value)

                if not match_number and not home_name and not away_name and not league_code:
                    skipped_count += 1
                    continue

                if not match_number or not home_name or not away_name or not league_code:
                    raise CommandError(
                        f"Row {row_number}: Missing required fixture data. "
                        f"MATCH NO='{match_number}', LEAGUE='{league_code}', "
                        f"HOME='{home_name}', AWAY='{away_name}'"
                    )

                league = self.get_league(league_code, row_number)

                try:
                    league_season = LeagueSeason.objects.get(
                        league=league,
                        season=season,
                    )
                except LeagueSeason.DoesNotExist:
                    raise CommandError(
                        f"Row {row_number}: LeagueSeason not found for "
                        f"{league} - {season}"
                    )

                home_team = self.get_team(home_name, row_number, "Home", league)
                away_team = self.get_team(away_name, row_number, "Away", league)

                match_date, status = self.normalise_date_and_status(match_date_raw, row_number)
                match_time = self.normalise_time(match_time_raw)

                match, created = Match.objects.update_or_create(
                    match_number=match_number,
                    defaults={
                        "league_season": league_season,
                        "home_team": home_team,
                        "away_team": away_team,
                        "match_date": match_date,
                        "match_time": match_time,
                        "venue": venue,
                        "status": status,
                    },
                )

                if created:
                    created_count += 1
                else:
                    updated_count += 1

        self.stdout.write(self.style.SUCCESS("Fixture import complete"))
        self.stdout.write(f"Created: {created_count}")
        self.stdout.write(f"Updated: {updated_count}")
        self.stdout.write(f"Skipped blank rows: {skipped_count}")